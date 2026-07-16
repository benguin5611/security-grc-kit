#!/usr/bin/env python3
"""Generate an ISO 27001 Statement of Applicability spreadsheet from a YAML control list."""

import argparse
import sys
from datetime import date

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

STATUS_FILLS = {
    "Implemented": "C6E2B5",
    "Partially implemented": "FFE699",
    "Planned": "FCD5B4",
    "Not implemented": "F4B7B7",
}
NOT_APPLICABLE_FILL = "D9D9D9"
HEADER_FILL = "1F3864"
HEADER_FONT_COLOUR = "FFFFFF"

COLUMNS = [
    ("Control ID", 12),
    ("Title", 42),
    ("Theme", 16),
    ("Applicable", 12),
    ("Justification", 40),
    ("Implementation Status", 20),
    ("Owner", 18),
    ("Notes", 40),
]


def load_controls(path):
    with open(path, "r", encoding="utf-8") as f:
        data = yaml.safe_load(f)
    if not data or "controls" not in data:
        raise ValueError(f"{path} has no top-level 'controls' list")
    return data


def set_cell(ws, row, column, value):
    """Write a cell, storing any '='-prefixed string as literal text.

    openpyxl treats a string value starting with '=' as a live Excel formula,
    so YAML fields (justification, notes, owner, ...) could otherwise inject
    executable formulas into the generated workbook.
    """
    cell = ws.cell(row=row, column=column, value=value)
    if isinstance(value, str) and value.startswith("="):
        cell.data_type = "s"
    return cell


def row_fill_for(control):
    if not control.get("applicable", True):
        return NOT_APPLICABLE_FILL
    return STATUS_FILLS.get(control.get("implementation_status", ""), None)


def write_controls_sheet(wb: Workbook, controls, org_name, as_of):
    ws = wb.active
    ws.title = "Statement of Applicability"
    set_cell(ws, 1, 1, f"{org_name}: Statement of Applicability (as of {as_of})")
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(COLUMNS))
    ws.cell(row=1, column=1).font = Font(bold=True, size=13)

    header_row = 3
    for col_idx, (title, width) in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_idx, value=title)
        cell.font = Font(bold=True, color=HEADER_FONT_COLOUR)
        cell.fill = PatternFill("solid", fgColor=HEADER_FILL)
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col_idx)].width = width
    ws.freeze_panes = f"A{header_row + 1}"

    row = header_row + 1
    for control in controls:
        fill = row_fill_for(control)
        values = [
            control.get("id", ""),
            control.get("title", ""),
            control.get("theme", ""),
            "Yes" if control.get("applicable", True) else "No",
            control.get("justification", ""),
            control.get("implementation_status", ""),
            control.get("owner", ""),
            control.get("notes", ""),
        ]
        for col_idx, value in enumerate(values, start=1):
            cell = set_cell(ws, row, col_idx, value)
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            if fill:
                cell.fill = PatternFill("solid", fgColor=fill)
        row += 1


def write_summary_sheet(wb: Workbook, controls):
    ws = wb.create_sheet("Summary")
    ws.cell(row=1, column=1, value="Summary").font = Font(bold=True, size=13)

    total = len(controls)
    applicable = sum(1 for c in controls if c.get("applicable", True))
    not_applicable = total - applicable

    ws.cell(row=3, column=1, value="Total controls")
    ws.cell(row=3, column=2, value=total)
    ws.cell(row=4, column=1, value="Applicable")
    ws.cell(row=4, column=2, value=applicable)
    ws.cell(row=5, column=1, value="Not applicable")
    ws.cell(row=5, column=2, value=not_applicable)

    ws.cell(row=7, column=1, value="By theme").font = Font(bold=True)
    themes = {}
    for c in controls:
        themes.setdefault(c.get("theme", "Unspecified"), 0)
        themes[c.get("theme", "Unspecified")] += 1
    r = 8
    for theme, count in sorted(themes.items()):
        set_cell(ws, r, 1, theme)
        ws.cell(row=r, column=2, value=count)
        r += 1

    r += 1
    ws.cell(row=r, column=1, value="By implementation status (applicable controls only)").font = Font(bold=True)
    r += 1
    statuses = {}
    for c in controls:
        if not c.get("applicable", True):
            continue
        status = c.get("implementation_status", "Unspecified") or "Unspecified"
        statuses.setdefault(status, 0)
        statuses[status] += 1
    for status, count in sorted(statuses.items()):
        set_cell(ws, r, 1, status)
        ws.cell(row=r, column=2, value=count)
        fill = STATUS_FILLS.get(status)
        if fill:
            ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor=fill)
        r += 1

    ws.column_dimensions["A"].width = 45
    ws.column_dimensions["B"].width = 12


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("input", help="Path to the YAML control list")
    parser.add_argument("-o", "--output", default="statement-of-applicability.xlsx", help="Output .xlsx path")
    args = parser.parse_args()

    data = load_controls(args.input)
    controls = data["controls"]
    org_name = data.get("organisation", "Your Organisation")
    as_of = data.get("date", str(date.today()))

    wb = Workbook()
    write_controls_sheet(wb, controls, org_name, as_of)
    write_summary_sheet(wb, controls)
    wb.save(args.output)
    print(f"Wrote {args.output} ({len(controls)} controls)")


if __name__ == "__main__":
    sys.exit(main())
